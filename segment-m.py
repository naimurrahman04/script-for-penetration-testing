import socket
import subprocess
import time
import os
import pyfiglet
from docx import Document
from docx.shared import Inches
from tqdm import tqdm


# Print the banner with big bold red text centered in the console
banner_text = "\033[1;31;40m\033[1;7m\033[6;1mTESTING STARTED\033[0m"
console_width = os.get_terminal_size().columns
centered_banner = banner_text.center(console_width)

ascii_banner = pyfiglet.figlet_format("EIC SEGMENTATION TEST")
print(ascii_banner)
print(centered_banner)

# Get local IP address
s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
s.connect(("8.8.8.8", 80))
local_ip_address = s.getsockname()[0]
s.close()

# Define input and output file paths
input_file_path = "input.txt"
output_file_path = "output.docx"

# Read IP addresses from input file
with open(input_file_path, "r") as input_file:
    ip_addresses = input_file.read().splitlines()

# Create Word document and add a table
document = Document()
table = document.add_table(rows=len(ip_addresses) + 1, cols=5)
table.style = "Table Grid"

# Set table headers
headers = ["Probe Source", "Destination VLAN", "Scan State", "Compliance Status", "Accessible/Not Accessible"]
for i, header in enumerate(headers):
    table.cell(0, i).text = header

# Ping each IP address and write results to table
ping_log = []
with tqdm(ip_addresses, desc="Pinging IPs", bar_format="\033[1;31m{l_bar}{bar}\033[0m") as progress_bar:
    for i, ip in enumerate(progress_bar):
        ping_process = subprocess.Popen(["ping", "-c", "3", ip], stdout=subprocess.PIPE, universal_newlines=True)
        output, _ = ping_process.communicate()
        ping_result = ping_process.returncode
        compliant = "non-compliant" if ping_result == 0 else "compliant"
        table.cell(i + 1, 0).text = local_ip_address + " (CDE In-Scope)"
        table.cell(i + 1, 1).text = ip
        table.cell(i + 1, 2).text = "CDE In-Scope to Out-of-Scope"
        table.cell(i + 1, 3).text = compliant
        table.cell(i + 1, 4).text = "Accessible" if compliant == "non-compliant" else "Not Accessible"
	
        # Store ping log
        ping_log.append(f"\033[1;33mPinging {ip}... \033[{'31m' if ping_result == 0 else '91m'}{compliant}\033[0m\n{output}")

        time.sleep(0.5)  # Wait for half a second before pinging the next IP address

# Save Word document
document.save(output_file_path)

# Print ping log
for log in ping_log:
    print(" ")
    print("\033[1;32m" + log + "\033[0m")

print(f"\n\033[1mResults saved to {output_file_path}\033[0m")
